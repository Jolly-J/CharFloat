#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
宣传片道具生成器（两套环境）。

设计约束：
1. **全部数据为 CLI 合成的虚构数据**，不含任何真实业务数据、真实客户名或真实凭据。
   演示数字沿用官网双轨动画口径：营收 86/94/108，成本 66/72/80，利润 20/22/28，利润率 23.3%/23.4%/25.9%。
2. 幂等：重复执行覆盖生成，录制把演示文件改坏了，重跑一次即可复原。
3. 只写入 promo/ 目录，不触碰仓库其他文件。

产物：
  env-A-无插件/01-发给AI的文件/月度经营.xlsx          源文件（格式完整，"已完成 80%"）
  env-A-无插件/02-AI返回的文件/*.xlsx                 5 个版本道具（摆拍：模拟云端 AI 重新生成的产物）
  env-A-无插件/03-加密文件/                            真·密码保护 xlsx + 加密 zip
  env-B-有插件/月度经营.xlsx                          与源文件同一份，录制时在 WPS 里打开

用法：
  python3 promo/assets/make-props.py
  python3 promo/assets/make-props.py --out /tmp/试造     # 试造，不覆盖正式道具
"""
from __future__ import annotations

import argparse
import shutil
import subprocess
import sys
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# 虚构数据（改这里即可换一套演示故事）
# ---------------------------------------------------------------------------
DEMO_PASSWORD = "Demo#2026"          # 演示专用口令，非真实凭据
TITLE = "2026 年 7–9 月经营分析（内部）"
SHEET_MAIN = "月度经营分析"
SHEET_DETAIL = "9月明细"
HEADERS = ["经营指标 / 万元", "7 月", "8 月", "9 月"]
INDICATORS = ["营业收入", "运营成本", "营业利润", "利润率"]
JUL = [86, 66]                        # 7 月：营收、成本
AUG = [94, 72]                        # 8 月：营收、成本
SEP = [108, 80]                       # 9 月：营收、成本（在明细表里，需汇总）
INCOME_ITEMS = [("华东区", 42), ("华南区", 31), ("华北区", 22), ("海外", 13)]
COST_ITEMS = [("采购", 46), ("人力", 21), ("渠道", 9), ("其他", 4)]
GROUP_UNITS = [                       # 加密道具用：虚构集团子公司
    ("星海科技", 4820), ("云岭实业", 3640), ("南屿数科", 2915), ("北辰智造", 2380),
    ("澄月传媒", 1740), ("长风物流", 1520), ("锦鲤新材", 1180), ("见山软件", 960),
    ("拾光文旅", 720), ("行远能源", 640), ("禾下农业", 380), ("青禾医疗", 260),
]

# 配色
C_TITLE_BG = "1F3864"     # 你的品牌深蓝
C_HEAD_BG = "2F5597"
C_HEAD_FG = "FFFFFF"
C_BORDER = "D6DCE4"
C_AI_BG = "ED7D31"        # 云端 AI 自己选的橙色——不是你的蓝
CHART_BAR = "3B82F6"
CHART_LINE = "0F766E"

FONT_CJK_CANDIDATES = [
    "/System/Library/Fonts/STHeiti Light.ttc",
    "/System/Library/Fonts/Hiragino Sans GB.ttc",
    "/System/Library/Fonts/Supplemental/Songti.ttc",
    "/Library/Fonts/Arial Unicode.ttf",
]


# ---------------------------------------------------------------------------
# 通用小工具
# ---------------------------------------------------------------------------
def thin_border() -> Border:
    side = Side(style="thin", color=C_BORDER)
    return Border(left=side, right=side, top=side, bottom=side)


def set_widths(ws, widths: dict[str, float]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def add_defined_name(wb, name: str, ref: str) -> bool:
    """openpyxl 3.1 改过 defined_names 的 API，两种写法都兼容。"""
    from openpyxl.workbook.defined_name import DefinedName
    dn = DefinedName(name, attr_text=ref)
    try:
        wb.defined_names.add(dn)          # openpyxl < 3.1
        return True
    except AttributeError:
        pass
    try:
        wb.defined_names[name] = dn       # openpyxl >= 3.1
        return True
    except Exception:
        return False


def cjk_font(size: int):
    from PIL import ImageFont
    for path in FONT_CJK_CANDIDATES:
        if Path(path).exists():
            try:
                return ImageFont.truetype(path, size)
            except Exception:
                continue
    return ImageFont.load_default()


# ---------------------------------------------------------------------------
# 主表内容
# ---------------------------------------------------------------------------
def write_main_sheet(ws, *, with_format: bool, sep_filled: bool, sheet_label: str = SHEET_MAIN,
                     ai_chart_image: Path | None = None, ai_palette: bool = False,
                     reorder: bool = False, extra_empty_sheet_row: str | None = None,
                     use_formulas: bool = False) -> None:
    """把演示数据写进工作表。

    with_format=False  → 模拟云端 AI 重新生成的产物：填色/边框/批注/冻结/数字格式全部丢失
    sep_filled=True    → 9 月列已并入（AI 返回的版本）
    reorder=True       → 指标行顺序被 AI 重排
    """
    rows = list(INDICATORS)
    if reorder:
        rows = ["营业收入", "营业利润", "利润率", "运营成本"]
    values = {
        "营业收入": [JUL[0], AUG[0], SEP[0]],
        "运营成本": [JUL[1], AUG[1], SEP[1]],
        "营业利润": [JUL[0] - JUL[1], AUG[0] - AUG[1], SEP[0] - SEP[1]],
        "利润率": [JUL[0] - JUL[1], AUG[0] - AUG[1], SEP[0] - SEP[1]],
    }

    if with_format:
        ws.merge_cells("A1:D1")
        ws["A1"] = TITLE
        ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
        ws["A1"].fill = PatternFill("solid", fgColor=C_TITLE_BG)
        ws["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[1].height = 32

    header_bg = C_AI_BG if ai_palette else C_HEAD_BG
    for idx, text in enumerate(HEADERS, start=1):
        cell = ws.cell(row=2, column=idx, value=text)
        cell.alignment = Alignment(horizontal="center" if idx > 1 else "left", vertical="center")
        if with_format:
            cell.font = Font(bold=True, size=11, color=C_HEAD_FG)
            cell.fill = PatternFill("solid", fgColor=header_bg)
    if with_format:
        ws.row_dimensions[2].height = 22

    # A1 标题占了第 1 行；无格式版本直接写裸标题，保持"这不是你那份文件"的观感
    if not with_format:
        ws["A1"] = TITLE

    ratio_fmt = "0.0%" if with_format else "General"
    money_fmt = "#,##0" if with_format else "General"
    for offset, name in enumerate(rows):
        r = 3 + offset
        label = ws.cell(row=r, column=1, value=name)
        if with_format:
            label.font = Font(bold=(name in ("营业利润", "利润率")))
        for col, raw in zip((2, 3, 4), values[name]):
            if col == 4 and not sep_filled:
                continue
            letter = get_column_letter(col)
            if name == "利润率":
                # 有插件版本的源文件用动态公式；云端 AI 重新生成的版本只剩死值
                expr = f"={letter}5/{letter}3"
                cell = ws.cell(row=r, column=col, value=expr if use_formulas else raw / values["营业收入"][col - 2])
                cell.number_format = ratio_fmt
            elif name == "营业利润":
                expr = f"={letter}3-{letter}4"
                cell = ws.cell(row=r, column=col, value=expr if use_formulas else raw)
                cell.number_format = money_fmt
            else:
                cell = ws.cell(row=r, column=col, value=raw)
                cell.number_format = money_fmt

    if with_format:
        r0, r1 = 3, 3 + len(rows) - 1
        for r in range(2, r1 + 1):
            for c in range(1, 5):
                ws.cell(row=r, column=c).border = thin_border()
        ws.cell(row=2, column=1).comment = None
        ws["C3"].comment = Comment("8 月数据已与财务对账一致，勿改。", "财务-林")
        ws["C3"].comment.width = 190
        ws["C3"].comment.height = 70
        ws.freeze_panes = "A3"
        set_widths(ws, {"A": 16.5, "B": 13.2, "C": 12.6, "D": 14.0})
        for r in range(3, r1 + 1):
            ws.row_dimensions[r].height = 20
    else:
        set_widths(ws, {"A": 21.0, "B": 19.4, "C": 19.4, "D": 19.4})

    if extra_empty_sheet_row:
        ws.cell(row=len(rows) + 5, column=1, value=extra_empty_sheet_row)

    if ai_chart_image and ai_chart_image.exists():
        img = XLImage(str(ai_chart_image))
        img.width, img.height = 460, 230
        ws.add_image(img, "F2")

    ws.title = sheet_label


def build_base_workbook(path: Path) -> None:
    """源文件：格式完整、9 月待补、9 月明细在第二个工作表。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    write_main_sheet(ws, with_format=True, sep_filled=False, sheet_label=SHEET_MAIN, use_formulas=True)

    det = wb.create_sheet(SHEET_DETAIL)
    det["A1"] = "9 月明细（分项，待并入分析表）"
    det["A1"].font = Font(bold=True, size=12)
    det["A2"] = "数据来源：各区域 9 月报表汇总，尚未计算合计"
    det["A2"].font = Font(size=9, color="808080")

    def block(start_row: int, title: str, items: list[tuple[str, int]]) -> None:
        head = det.cell(row=start_row, column=1, value=title)
        head.font = Font(bold=True, color="FFFFFF")
        head.fill = PatternFill("solid", fgColor=C_HEAD_BG)
        det.cell(row=start_row, column=2, value="金额（万元）").font = Font(bold=True, color="FFFFFF")
        det.cell(row=start_row, column=2).fill = PatternFill("solid", fgColor=C_HEAD_BG)
        for i, (name, amount) in enumerate(items, start=1):
            det.cell(row=start_row + i, column=1, value=name)
            det.cell(row=start_row + i, column=2, value=amount).number_format = "#,##0"
        total_row = start_row + len(items) + 1
        det.cell(row=total_row, column=1, value="合计").font = Font(bold=True)
        total = det.cell(row=total_row, column=2)   # 故意留空：等 AI 算
        total.fill = PatternFill("solid", fgColor="FFF2CC")   # 用户自己标的"待处理"浅黄
        total.border = thin_border()

    block(4, "收入分项", INCOME_ITEMS)
    block(11, "成本分项", COST_ITEMS)
    set_widths(det, {"A": 18.0, "B": 14.0})
    det.freeze_panes = "A4"

    add_defined_name(wb, "经营数据", f"'{SHEET_MAIN}'!$B$3:$D$6")
    wb.save(path)


# ---------------------------------------------------------------------------
# 图片版"图表"（v2 道具：AI 给的图是张图片，不能在表里改数据）
# ---------------------------------------------------------------------------
def draw_chart_png(path: Path) -> None:
    from PIL import Image, ImageDraw

    W, H = 920, 460
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)
    f_title = cjk_font(30)
    f_axis = cjk_font(22)
    f_val = cjk_font(24)

    d.text((30, 18), "营收与利润率趋势", font=f_title, fill="#1f2937")
    left, right, top, bottom = 90, W - 60, 110, H - 80
    for i in range(5):
        y = top + (bottom - top) * i / 4
        d.line([(left, y), (right, y)], fill="#E8EEF6", width=2)
    d.line([(left, bottom), (right, bottom)], fill="#94a3b8", width=2)

    months = ["7 月", "8 月", "9 月"]
    revenue = [86, 94, 108]
    ratio = [23.3, 23.4, 25.9]
    slot = (right - left) / 3
    max_rev = 120.0
    for i, (m, rev) in enumerate(zip(months, revenue)):
        cx = left + slot * (i + 0.5)
        bar_h = (bottom - top) * rev / max_rev
        shade = ["#BFDBFE", "#93C5FD", "#3B82F6"][i]
        d.rectangle([cx - 45, bottom - bar_h, cx + 45, bottom], fill=shade)
        d.text((cx - 26, bottom - bar_h - 32), str(rev), font=f_val, fill="#2563eb")
        d.text((cx - 22, bottom + 14), m, font=f_axis, fill="#64748b")

    pts = []
    for i, r in enumerate(ratio):
        cx = left + slot * (i + 0.5)
        y = top + (bottom - top) * (30 - r) / 30
        pts.append((cx, y))
    d.line(pts, fill="#0f766e", width=4, joint="curve")
    for (x, y), r in zip(pts, ratio):
        d.ellipse([x - 7, y - 7, x + 7, y + 7], fill="white", outline="#0f766e", width=4)
        d.text((x - 24, y - 40), f"{r}%", font=f_axis, fill="#0f766e")

    d.text((left - 62, top - 34), "万元", font=f_axis, fill="#94a3b8")
    d.text((right - 10, top - 34), "%", font=f_axis, fill="#94a3b8")
    img.save(path, "PNG")


# ---------------------------------------------------------------------------
# AI 返回的 5 个版本道具（摆拍：模拟"每次都要重新生成一份"）
# ---------------------------------------------------------------------------
def build_ai_returned(out_dir: Path, chart_png: Path) -> list[str]:
    made: list[str] = []

    def save(wb, filename: str) -> None:
        wb.save(out_dir / filename)
        made.append(filename)

    # v1：分析做完了，但你原有的格式全丢了，公式变成死值
    wb = openpyxl.Workbook()
    write_main_sheet(wb.active, with_format=False, sep_filled=True, sheet_label="Sheet1")
    save(wb, "月度经营_分析报表.xlsx")

    # v2：你说"格式没了"，它按自己的审美重新上色（橙色），图表是张图片
    wb = openpyxl.Workbook()
    write_main_sheet(wb.active, with_format=True, sep_filled=True, sheet_label="Sheet1",
                     ai_chart_image=chart_png, ai_palette=True)
    save(wb, "月度经营_分析报表_v2.xlsx")

    # v2_最新：同一句话第二次执行，结果又不一样（指标顺序被重排）
    wb = openpyxl.Workbook()
    write_main_sheet(wb.active, with_format=True, sep_filled=True, sheet_label="Sheet1",
                     ai_chart_image=chart_png, ai_palette=True, reorder=True)
    save(wb, "月度经营_分析报表_v2_最新.xlsx")

    # 最终版：数字齐了，但多出一个空工作表
    wb = openpyxl.Workbook()
    write_main_sheet(wb.active, with_format=True, sep_filled=True, sheet_label="Sheet1")
    wb.create_sheet("Sheet2")
    save(wb, "月度经营_分析报表_最终版.xlsx")

    # 真的最终版：连标题都被它改写成了"AI 生成"
    wb = openpyxl.Workbook()
    ws = wb.active
    write_main_sheet(ws, with_format=True, sep_filled=True, sheet_label="Sheet1")
    ws["A1"] = "月度经营分析报表（AI 生成）"
    ws.cell(row=9, column=1, value="* 本报表由 AI 重新生成，原始格式与批注未保留").font = Font(size=9, color="C00000")
    save(wb, "月度经营_分析报表_最终版_真的最终版.xlsx")

    return made


# ---------------------------------------------------------------------------
# 加密道具
# ---------------------------------------------------------------------------
def build_encrypted(out_dir: Path, stage: Path) -> dict[str, str]:
    """真·密码保护 xlsx（ECMA-376 标准加密，Excel/WPS 打开需口令）+ 加密 zip 兜底。

    明文底稿只存在于临时目录，正式目录里不留（避免"机密文件旁边就放着明文"的穿帮）。
    """
    plain = stage / "2026年度集团经营数据.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "集团合并口径"
    ws.merge_cells("A1:C1")
    ws["A1"] = "2026 年度集团经营数据（密级：内部）"
    ws["A1"].font = Font(bold=True, size=13, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="7F1D1D")
    ws.row_dimensions[1].height = 28
    for i, text in enumerate(["子公司", "合并营收（万元）", "内部编号"], start=1):
        cell = ws.cell(row=2, column=i, value=text)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="9C2A2A")
    for i, (name, amount) in enumerate(GROUP_UNITS, start=3):
        ws.cell(row=i, column=1, value=name)
        ws.cell(row=i, column=2, value=amount).number_format = "#,##0"
        ws.cell(row=i, column=3, value=f"GF-{i - 2:03d}-X")
    set_widths(ws, {"A": 18.0, "B": 20.0, "C": 14.0})
    plain_row = 3 + len(GROUP_UNITS)
    ws.cell(row=plain_row, column=1, value="合计").font = Font(bold=True)
    ws.cell(row=plain_row, column=2, value=f"=SUM(B3:B{plain_row - 1})").number_format = "#,##0"
    wb.save(plain)

    enc_xlsx = out_dir / "2026年度集团经营数据_机密.xlsx"
    from msoffcrypto.format.ooxml import OOXMLFile
    with open(plain, "rb") as fin, open(enc_xlsx, "wb") as fout:
        OOXMLFile(fin).encrypt(DEMO_PASSWORD, fout)

    enc_zip = out_dir / "2026年度集团经营数据_机密.zip"
    if enc_zip.exists():
        enc_zip.unlink()
    subprocess.run(
        ["zip", "-q", "-j", "-P", DEMO_PASSWORD, str(enc_zip), str(plain)],
        check=True, capture_output=True, text=True,
    )
    return {
        "加密xlsx": enc_xlsx.name,
        "加密zip": enc_zip.name,
        "口令": DEMO_PASSWORD,
    }


# ---------------------------------------------------------------------------
def main() -> int:
    here = Path(__file__).resolve().parent
    promo = here.parent
    ap = argparse.ArgumentParser(description="生成宣传片两套环境道具")
    ap.add_argument("--out", default=str(promo), help="输出根目录（默认 promo/）")
    args = ap.parse_args()

    root = Path(args.out).resolve()
    env_a = root / "env-A-无插件"
    dir_src = env_a / "01-发给AI的文件"
    dir_ret = env_a / "02-AI返回的文件"
    dir_enc = env_a / "03-加密文件"
    env_b = root / "env-B-有插件"
    tmp = root / "assets" / ".build"
    for d in (dir_src, dir_ret, dir_enc, env_b, tmp):
        d.mkdir(parents=True, exist_ok=True)

    print("== 造源文件 ==")
    base = dir_src / "月度经营.xlsx"
    build_base_workbook(base)
    print(f"  {base.relative_to(root)}")

    print("== 造有插件环境（同一份文件的拷贝）==")
    demo = env_b / "月度经营.xlsx"
    shutil.copyfile(base, demo)
    print(f"  {demo.relative_to(root)}")

    print("== 造 AI 返回版本道具 ==")
    chart_png = tmp / "ai_chart_picture.png"
    draw_chart_png(chart_png)
    for name in build_ai_returned(dir_ret, chart_png):
        print(f"  {name}")

    print("== 造加密道具 ==")
    info = build_encrypted(dir_enc, tmp)
    for k, v in info.items():
        print(f"  {k}: {v}")

    # 明文底稿不进正式目录（避免"机密文件旁边就放着明文"这种穿帮）
    shutil.rmtree(tmp, ignore_errors=True)
    print("\n完成。数据全部为 CLI 合成的虚构数据。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
